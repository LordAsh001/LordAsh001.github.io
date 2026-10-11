"""One collector per source. Each `parse_*` is a pure function (tested in tests/),
each `collect_*` fetches the live page and hands it to the parser."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, timedelta

from bs4 import BeautifulSoup

from common import (get, kg_in_unit, match_commodity, match_market, norm_state, record, to_number, today)

MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


# ---------------------------------------------------------------- Kasuwa (Agriarche)
KASUWA_API = "https://agriarche-backend.onrender.com"


def parse_kasuwa_prices(rows: list[dict], max_age_days: int = 150) -> list[dict]:
    """Agent price records: start_time, state, market, commodity, price_per_bag, weight_of_bag_kg, price_per_kg."""
    out, cutoff = [], (date.today() - timedelta(days=max_age_days)).isoformat()
    for r in rows:
        c = match_commodity(r.get("commodity", ""))
        d = (r.get("start_time") or "")[:10]
        if not c or not d or d < cutoff:
            continue
        st = norm_state(r.get("state"))
        mid = match_market(r.get("market", ""), st)
        pk, bag, kg = to_number(r.get("price_per_kg")), to_number(r.get("price_per_bag")), to_number(r.get("weight_of_bag_kg"))
        if not pk and bag and kg:
            pk = bag / kg
        if not pk:
            continue
        unit = f"1 kg ({int(kg)} kg bag at ₦{int(bag):,})" if bag and kg else "1 kg"
        if mid:
            out.append(record("kasuwa_api", c, pk, unit, 1, d, "market", [mid], name=r.get("market"), state=st,
                              availability=r.get("availability")))
        elif st:
            out.append(record("kasuwa_api", c, pk, unit, 1, d, "state", st, name=r.get("market")))
    return out


def parse_kasuwa_open(rows: list[dict]) -> list[dict]:
    """Open-market records: date, commodity, location ('Giwa Market, Kaduna State'), unit, price (per bag, size not given)."""
    out = []
    for r in rows:
        c = match_commodity(r.get("commodity", ""))
        loc = r.get("location") or ""
        parts = [p.strip() for p in loc.split(",")]
        st = norm_state(parts[-1]) if len(parts) > 1 else None
        mid = match_market(parts[0], st)
        price = to_number(r.get("price"))
        d = (r.get("date") or "")[:10]
        if not (c and price and d and (mid or st)):
            continue
        per = kg_in_unit(r.get("unit") or "")
        out.append(record("kasuwa_api", c, price, r.get("unit") or "bag (size not stated)", per, d,
                          "market" if mid else "state", [mid] if mid else st, name=parts[0], kind="open-market"))
    return out


def collect_kasuwa():
    rows = []
    for page in range(1, 8):
        j = get(f"{KASUWA_API}/prices/filtered", as_json=True,
                params={"commodity": "", "market": "", "page": page, "page_size": 100})
        data = j.get("data") or []
        rows += data
        if len(data) < 100 or (data and data[-1].get("start_time", "") < (date.today() - timedelta(days=150)).isoformat()):
            break
    recs = parse_kasuwa_prices(rows)
    j = get(f"{KASUWA_API}/other-sources/filtered", as_json=True,
            params={"page": 1, "page_size": 300, "commodity": "", "location": ""})
    recs += parse_kasuwa_open(j.get("data") or [])
    return recs


# ---------------------------------------------------------------- Lagos State food price tracker (xlsx)
LAGOS_PAGE = "https://lagosagric.com/price-tracker/"
LAGOS_MARKET_SHEETS = {
    "Mushin Market": "mushin", "Mile 12 Market": "mile12", "Ile-Epo Market": "ileepo", "Sabo Market": "sabo",
    "Iddo Oyingbo Market": "oyingbo", "Agege Market": "agege", "Ajah Market": "ajah", "Agbalata Market": "agbalata",
    "Fresh Food Hub - Mushin": "ffhmushin", "Epe Market": "epe",
}
# (pattern on the item name, commodity key, base units in the pack)
LAGOS_ITEMS = [
    (r"^50kg bag of nigerian", "ricel", 50), (r"^50kg bag of imported rice \(long", "ricei", 50),
    (r"3\.3kg .*oloyin", "oloyin", 3.3), (r"3\.3kg .*drum", "beansb", 3.3),
    (r"2\.6kg .*garri \(white", "garriw", 2.6), (r"2\.6kg .*garri \(yellow", "garriy", 2.6),
    (r"crate of medium sized eggs", "eggs", 1), (r"tuber of medium sized yam", "yam", 1),
    (r"^1kg of tomato", "tomato", 1), (r"^1kg of pepper", "pepper", 1), (r"^1kg of onion", "onion", 1),
]


def lagos_month(label: str, fallback_year: int) -> str | None:
    t = label.lower()
    m = re.search(r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*'?\s*(\d{2,4})?", t)
    if not m:
        return None
    y = m.group(2)
    year = fallback_year if not y else (2000 + int(y) if len(y) == 2 else int(y))
    return f"{year}-{MONTHS[m.group(1)]:02d}"


def parse_lagos_workbook(wb, keep_months: int = 3) -> list[dict]:
    out, series = [], {}
    for sheet, mid in LAGOS_MARKET_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(rows) if len(r) > 1 and isinstance(r[1], str) and r[1].strip().upper() == "FOOD ITEM"), None)
        if hi is None:
            continue
        hdr = rows[hi]
        avg_cols = [(i, lagos_month(str(h), date.today().year)) for i, h in enumerate(hdr) if isinstance(h, str) and "average" in h.lower()]
        avg_cols = [(i, m) for i, m in avg_cols if m]
        for r in rows[hi + 1:]:
            if len(r) < 2 or not isinstance(r[1], str):
                continue
            name = r[1].strip().lower()
            spec = next(((c, per) for pat, c, per in LAGOS_ITEMS if re.search(pat, name)), None)
            if not spec:
                continue
            c, per = spec
            vals = [(m, r[i]) for i, m in avg_cols if i < len(r) and isinstance(r[i], (int, float)) and r[i] > 0]
            for m, v in vals:
                series.setdefault((c, per, r[1].strip(), m), []).append(v)
            for m, v in vals[-keep_months:]:
                out.append(record("lagos_tracker", c, v, r[1].strip(), per, m, "market", [mid]))
    for (c, per, unit, m), vs in series.items():
        out.append(record("lagos_tracker", c, sum(vs) / len(vs), unit, per, m, "state", "Lagos", series="lagos10"))
    return out


def collect_lagos():
    from openpyxl import load_workbook
    html = get(LAGOS_PAGE)
    links = re.findall(r'href="([^"]+\.xlsx)"', html)
    if not links:
        raise RuntimeError("no .xlsx link on the tracker page")
    xb = get(links[0], binary=True)
    wb = load_workbook(io.BytesIO(xb), read_only=True, data_only=True)
    return parse_lagos_workbook(wb)


# ---------------------------------------------------------------- LCFE market data
LCFE_URL = "https://lcfe.ng/market-data.php"


def parse_lcfe(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for table in soup.find_all("table"):
        heads = [th.get_text(strip=True).upper() for th in table.find_all("th")]
        if "COMMODITIES" not in heads or "UNIT PRICE" not in heads:
            continue
        ix = {h: i for i, h in enumerate(heads)}
        for tr in table.find_all("tr"):
            tds = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
            if len(tds) < len(heads):
                continue
            c = match_commodity(tds[ix["COMMODITIES"]])
            price = to_number(tds[ix["UNIT PRICE"]])
            try:
                d = datetime.strptime(tds[ix["DATE"]], "%d/%m/%Y").date().isoformat()
            except ValueError:
                continue
            if not (c and price):
                continue
            unit = tds[ix["UNIT"]]
            per = kg_in_unit(unit) or 1
            for loc in re.split(r"[/,]", tds[ix.get("LOCATION", 0)]):
                st = norm_state(loc)
                if st:
                    out.append(record("lcfe", c, price, unit.lower(), per, d, "state", st))
    return out


def collect_lcfe():
    return parse_lcfe(get(LCFE_URL))


# ---------------------------------------------------------------- NCX ticker
NCX_URL = "https://ncx.com.ng/our-commodities/agricultural-commodities/"
NCX_STATE = {"PLT": "Plateau", "KAN": "Kano", "KAD": "Kaduna", "SKT": "Sokoto", "JIG": "Jigawa", "EBO": "Ebonyi",
             "TRB": "Taraba", "NAS": "Nasarawa", "YBE": "Yobe", "OND": "Ondo", "ND": "Ondo", "BEN": "Benue",
             "NIG": "Niger", "KAT": "Katsina", "ZAM": "Zamfara", "BAU": "Bauchi", "GOM": "Gombe", "KEB": "Kebbi",
             "KWA": "Kwara", "ABJ": "FCT", "FCT": "FCT", "ADA": "Adamawa", "OYO": "Oyo"}


def parse_ncx(html: str) -> list[dict]:
    text = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)
    out, seen = [], set()
    for name, code, price in re.findall(r"([A-Za-z][A-Za-z ]{2,30}?)\s*\(([A-Z]{5,7})\)\s*₦\s*([\d,]+(?:\.\d+)?)", text):
        if code in seen:
            continue
        seen.add(code)
        c = match_commodity(name)
        st = NCX_STATE.get(code[3:]) or NCX_STATE.get(code[-2:])
        p = to_number(price)
        if c and st and p:
            out.append(record("ncx_ticker", c, p, "1 kg (inferred)", 1, today(), "state", st, undated=True, code=code))
    return out


def collect_ncx():
    return parse_ncx(get(NCX_URL))


# ---------------------------------------------------------------- Commodity Nigeria (national 50 kg board)
CNG_URL = "https://commodity.ng/live-prices/"


def parse_commodity_ng(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for table in soup.find_all("table"):
        heads = [th.get_text(" ", strip=True) for th in table.find_all("th")]
        col = next((i for i, h in enumerate(heads) if re.search(r"50\s*kg", h, re.I)), None)
        if col is None:
            continue
        for tr in table.find_all("tr")[1:]:
            tds = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
            if len(tds) <= col:
                continue
            c, p = match_commodity(tds[0]), to_number(tds[col])
            if c and p:
                out.append(record("commodity_ng", c, p, "50 kg bag", 50, today(), "national", "Nigeria", undated=True))
    return out


def collect_commodity_ng():
    return parse_commodity_ng(get(CNG_URL))


# ---------------------------------------------------------------- Avestack wholesale index
AVESTACK_URL = "https://avestack.com/resources/market-reports/"


def parse_avestack(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for table in soup.find_all("table"):
        heads = [th.get_text(" ", strip=True).upper() for th in table.find_all("th")]
        if "AVERAGE PRICE" not in heads:
            continue
        ix = {h: i for i, h in enumerate(heads)}
        for tr in table.find_all("tr")[1:]:
            tds = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
            if len(tds) < len(heads):
                continue
            c = match_commodity(tds[ix["FOOD ITEM"]])
            p = to_number(tds[ix["AVERAGE PRICE"]])
            unit = tds[ix["WHOLESALE UNIT"]]
            loc = tds[ix.get("MAIN MARKET LOCATION", 0)]
            m = re.match(r"(.+?)\s*\(([^)]+)\)", loc)
            st = norm_state(m.group(2)) if m else None
            mid = match_market(m.group(1), st) if m else None
            per = 100 if re.search(r"stack of 100", unit, re.I) else kg_in_unit(unit)
            if c and p and (mid or st):
                out.append(record("avestack", c, p, unit, per, today(), "market" if mid else "state",
                                  [mid] if mid else st, undated=True, name=loc))
    return out


def collect_avestack():
    return parse_avestack(get(AVESTACK_URL))


# ---------------------------------------------------------------- Opaindex (NBS monthly series)
OPA_URL = "https://opaindex.com/reports/agriculture-prices-nigeria"


def parse_opaindex(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for tr in soup.find_all("tr"):
        tds = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
        if len(tds) < 3:
            continue
        c = match_commodity(tds[0])
        price = next((to_number(t) for t in tds[1:] if "₦" in t), None)
        dm = next((re.search(r"(\d{1,2})\s+([A-Za-z]{3})[a-z]*\s+(\d{4})", t) for t in tds if re.search(r"\d{4}", t)), None)
        unit = next((t for t in tds[1:] if re.search(r"per|crate|kg", t, re.I) and "₦" not in t), "1 kg")
        if not (c and price and dm):
            continue
        month = f"{dm.group(3)}-{MONTHS[dm.group(2).lower()[:3]]:02d}"
        per = 1 if c == "eggs" else (kg_in_unit(unit) or 1)
        out.append(record("opaindex", c, price, unit, per, month, "national", "Nigeria"))
    return out


def collect_opaindex():
    return parse_opaindex(get(OPA_URL))


# ---------------------------------------------------------------- WFP food prices on HDX (open CSV, CC BY-IGO)
HDX_API = "https://data.humdata.org/api/3/action/package_show"


def parse_wfp_csv(text: str, months_back: int = 6) -> list[dict]:
    cutoff = (date.today() - timedelta(days=31 * months_back)).isoformat()
    out = []
    rd = csv.DictReader(io.StringIO(text))
    for r in rd:
        d = (r.get("date") or "")[:10]
        if not d or d.startswith("#") or d < cutoff or (r.get("currency") or "NGN") != "NGN":
            continue
        c = match_commodity(r.get("commodity", ""))
        p = to_number(r.get("price"))
        if not (c and p):
            continue
        st = norm_state(r.get("admin1"))
        mid = match_market(r.get("market", ""), st)
        per = kg_in_unit(r.get("unit", "")) or (1 if (r.get("unit") or "").upper() == "KG" else None)
        out.append(record("wfp_hdx", c, p, (r.get("unit") or "").lower(), per, d[:7], "market" if mid else "state",
                          [mid] if mid else st, name=r.get("market"), pricetype=r.get("pricetype")))
    return [o for o in out if o["where"]]


def collect_wfp():
    meta = get(HDX_API, as_json=True, params={"id": "wfp-food-prices-for-nigeria"})
    res = next(r for r in meta["result"]["resources"] if r.get("format", "").upper() == "CSV" and "qc" not in r.get("name", "").lower())
    return parse_wfp_csv(get(res["url"]))


REGISTRY = {
    "kasuwa_api": collect_kasuwa,
    "lagos_tracker": collect_lagos,
    "lcfe": collect_lcfe,
    "ncx_ticker": collect_ncx,
    "commodity_ng": collect_commodity_ng,
    "avestack": collect_avestack,
    "opaindex": collect_opaindex,
    "wfp_hdx": collect_wfp,
}
